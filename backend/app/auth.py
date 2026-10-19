import hashlib,re,secrets,time,threading
from collections import defaultdict,deque
from datetime import timedelta
from typing import Annotated,Literal
from fastapi import APIRouter,Depends,Request,Response,HTTPException,UploadFile,File
from pydantic import BaseModel,Field,field_validator
from sqlalchemy import select,delete
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from argon2 import PasswordHasher
from argon2.exceptions import VerificationError,InvalidHashError
from .models import User,AuthSession,now

router=APIRouter(prefix='/api')
hasher=PasswordHasher(time_cost=2,memory_cost=32768,parallelism=1)
def database(request:Request):
    with request.app.state.db() as db:yield db
DB=Annotated[Session,Depends(database)]
def fail(code,message):raise HTTPException(code,message)
def verify(encoded,password):
    try:return hasher.verify(encoded,password)
    except (VerificationError,InvalidHashError):return False
def user_json(u):return {'id':u.id,'username':u.username,'name':u.name,'class_name':u.class_name,'role':u.role,'active':u.active,'must_change_password':u.must_change_password}
def get_user(request:Request,db:DB):
    token=request.cookies.get('lumaread_session','')
    session=db.get(AuthSession,hashlib.sha256(token.encode()).hexdigest()) if token else None
    if not session or session.expires_at<now():fail(401,'请先登录')
    user=db.get(User,session.user_id)
    if not user or not user.active:fail(401,'账号暂不可用，请联系机构')
    if user.must_change_password and request.url.path not in ['/api/auth/me','/api/auth/password','/api/auth/logout']:fail(403,'请先修改临时密码')
    return user
Current=Annotated[User,Depends(get_user)]
def require_admin(user:Current):
    if user.role!='admin':fail(403,'没有管理权限')
    return user
Admin=Annotated[User,Depends(require_admin)]

class Login(BaseModel):
    username:str=Field(min_length=1,max_length=32)
    password:str=Field(min_length=1,max_length=128)
class PasswordChange(BaseModel):
    current_password:str=Field(max_length=128)
    new_password:str=Field(min_length=8,max_length=64)
    @field_validator('new_password')
    @classmethod
    def password_rules(cls,v):
        if not re.search('[A-Za-z]',v) or not re.search('[0-9]',v):raise ValueError('密码需要字母和数字')
        return v
class NewUser(BaseModel):
    username:str=Field(pattern=r'^[a-zA-Z0-9_]{3,32}$')
    name:str=Field(min_length=1,max_length=30)
    class_name:str=Field(default='',max_length=60)
    password:str|None=Field(default=None,min_length=8,max_length=64)
    @field_validator('password')
    @classmethod
    def password_rules(cls,v):return PasswordChange.password_rules(v) if v else v
class UserChange(BaseModel):
    active:bool|None=None
    reset_password:bool=False
    name:str|None=Field(default=None,min_length=1,max_length=30)

def set_session(request,db,response,user):
    token=secrets.token_urlsafe(40)
    db.add(AuthSession(token_hash=hashlib.sha256(token.encode()).hexdigest(),user_id=user.id,expires_at=now()+timedelta(days=7)))
    response.set_cookie('lumaread_session',token,max_age=604800,httponly=True,secure=request.app.state.settings.secure_cookies,samesite='lax',path='/')

@router.post('/auth/login')
def login(data:Login,request:Request,response:Response,db:DB):
    key=(request.client.host,data.username.lower())
    with request.app.state.login_lock:
        attempts=request.app.state.login_attempts;cutoff=time.monotonic()-900
        for stale,q in list(attempts.items()):
            if not q or q[-1]<cutoff:del attempts[stale]
        if key not in attempts and len(attempts)>=4096:fail(429,'登录请求较多，请稍后再试')
        queue=attempts[key]
        while queue and queue[0]<cutoff:queue.popleft()
        if len(queue)>=10:fail(429,'尝试次数过多，请 15 分钟后再试')
        queue.append(time.monotonic())
    u=db.scalar(select(User).where(User.username==data.username.lower()))
    # Always verify a hash to avoid a fast missing-account branch.
    valid=verify(u.password_hash if u else request.app.state.dummy_hash,data.password)
    if not u or not valid or not u.active:fail(401,'账号或密码不正确，或账号已停用')
    old=request.cookies.get('lumaread_session')
    if old:db.execute(delete(AuthSession).where(AuthSession.token_hash==hashlib.sha256(old.encode()).hexdigest()))
    set_session(request,db,response,u);db.commit()
    with request.app.state.login_lock:request.app.state.login_attempts.pop(key,None)
    return user_json(u)

@router.get('/auth/me')
def me(user:Current):return user_json(user)

@router.post('/auth/logout')
def logout(request:Request,response:Response,db:DB):
    db.execute(delete(AuthSession).where(AuthSession.token_hash==hashlib.sha256(request.cookies.get('lumaread_session','').encode()).hexdigest()));db.commit()
    response.delete_cookie('lumaread_session',path='/');return {'ok':True}

@router.post('/auth/password')
def password(data:PasswordChange,request:Request,response:Response,user:Current,db:DB):
    if not verify(user.password_hash,data.current_password):fail(400,'当前密码不正确')
    if data.current_password==data.new_password:fail(400,'新密码不能与原密码相同')
    user.password_hash=hasher.hash(data.new_password);user.must_change_password=False
    db.execute(delete(AuthSession).where(AuthSession.user_id==user.id))
    set_session(request,db,response,user);db.commit();return user_json(user)

@router.get('/admin/users')
def users(admin:Admin,db:DB):return [user_json(u) for u in db.scalars(select(User).where(User.role=='student').order_by(User.created_at.desc()))]

@router.post('/admin/users',status_code=201)
def create_user(data:NewUser,admin:Admin,db:DB):
    password=data.password or secrets.token_urlsafe(10)+'9a'
    user=User(username=data.username.lower(),name=data.name,class_name=data.class_name,password_hash=hasher.hash(password))
    db.add(user)
    try:db.commit()
    except IntegrityError:db.rollback();fail(409,'登录账号已存在')
    return {**user_json(user),'temporary_password':password}

@router.patch('/admin/users/{user_id}')
def update_user(user_id:str,data:UserChange,admin:Admin,db:DB):
    u=db.get(User,user_id)
    if not u or u.role!='student':fail(404,'学生不存在')
    result={}
    if data.name is not None:u.name=data.name
    if data.active is not None:u.active=data.active
    if data.reset_password:
        pw=secrets.token_urlsafe(10)+'9a';u.password_hash=hasher.hash(pw);u.must_change_password=True;result['temporary_password']=pw
    if data.reset_password or data.active is False:db.execute(delete(AuthSession).where(AuthSession.user_id==u.id))
    db.commit();return {**user_json(u),**result}

def seed_admin(factory,username,password):
    with factory() as db:
        if db.scalar(select(User).where(User.username==username)):return
        db.add(User(username=username,name='管理员',password_hash=hasher.hash(password),role='admin',must_change_password=False));db.commit()


class StudentRow(BaseModel):
    name:str=Field(min_length=1,max_length=30)
    username:str=Field(default='',pattern=r'^(?:[a-zA-Z0-9_]{3,32})?$')
    class_name:str=Field(default='',max_length=60)
class StudentImport(BaseModel):
    rows:list[StudentRow]=Field(min_length=1,max_length=200)

@router.post('/admin/students/import',status_code=201)
def import_students(data:StudentImport,admin:Admin,db:DB):
    names=set();prepared=[]
    for row in data.rows:
        username=row.username.lower() or 's_'+secrets.token_hex(6)
        if not row.name.strip():fail(422,'姓名不能为空')
        if username in names or db.scalar(select(User.id).where(User.username==username)):fail(409,'账号重复：'+username)
        names.add(username);prepared.append((row,username))
    result=[]
    for row,username in prepared:
        pw=secrets.token_urlsafe(10)+'9a'
        u=User(username=username,name=row.name.strip(),class_name=row.class_name,password_hash=hasher.hash(pw));db.add(u);db.flush()
        result.append({**user_json(u),'temporary_password':pw})
    try:db.commit()
    except IntegrityError:db.rollback();fail(409,'账号已存在，请刷新后重试')
    return result

class StudentBulk(BaseModel):
    ids:list[str]=Field(min_length=1,max_length=200)
    action:Literal['enable','disable','reset','class']
    class_name:str=Field(default='',max_length=60)

@router.post('/admin/students/bulk')
def bulk_students(data:StudentBulk,admin:Admin,db:DB):
    users=list(db.scalars(select(User).where(User.id.in_(data.ids),User.role=='student').with_for_update()))
    if len(users)!=len(set(data.ids)):fail(422,'选择包含无效学生账号')
    result=[]
    for u in users:
        extra={}
        if data.action=='class':u.class_name=data.class_name
        if data.action in ['enable','disable']:u.active=data.action=='enable'
        if data.action=='reset':
            pw=secrets.token_urlsafe(10)+'9a';u.password_hash=hasher.hash(pw);u.must_change_password=True;extra['temporary_password']=pw
        if data.action in ['disable','reset']:db.execute(delete(AuthSession).where(AuthSession.user_id==u.id))
        result.append({**user_json(u),**extra})
    db.commit();return result

@router.post('/admin/students/preview')
async def preview_students(admin:Admin,file:UploadFile=File()):
    import io,zipfile
    from openpyxl import load_workbook
    raw=await file.read(2*1024*1024+1)
    if len(raw)>2*1024*1024:fail(422,'名单文件不能超过 2 MB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>20*1024*1024:fail(422,'表格内容过大')
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        rows=[]
        try:
            for i,values in enumerate(workbook.active.iter_rows(max_row=202,max_col=3,values_only=True)):
                cells=[str(v).strip() if v is not None else '' for v in values]
                if not any(cells) or (i==0 and cells[0]=='姓名'):continue
                if any(v.startswith('=') for v in cells):fail(422,'请将公式转换为普通文本后导入')
                rows.append(dict(zip(['name','username','class_name'],cells)))
        finally:workbook.close()
        if len(rows)>200:fail(422,'每批最多 200 人')
        return rows
    except HTTPException:raise
    except Exception:fail(422,'无法读取 Excel，请使用未加密的 .xlsx 文件')
